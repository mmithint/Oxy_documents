"""
HAZOP Report Exporter — generates Excel file matching OOG report format.
Uses openpyxl (already in requirements.txt).
"""
import io
from collections import defaultdict
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from app.models.hazop_models import HAZOPReport, Deviation

# Standard deviation type ordering (matches STANDARD_DEVIATION_TYPES)
DEVIATION_ORDER = [
    "High Pressure", "Low Pressure",
    "High Level", "Low Level",
    "High Temperature", "Low Temperature",
    "No/Low Flow", "More/High Flow",
    "Reverse / Misdirected Flow",
    "Tube Leak", "Composition / Contamination",
    "Human Factors", "Previous Incidents / Learnings", "Other",
]

HEADER_FILL  = PatternFill("solid", fgColor="1F3864")  # dark blue
SUBHEAD_FILL = PatternFill("solid", fgColor="2E75B6")  # medium blue
THIN_BORDER  = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"),  bottom=Side(style="thin"),
)

# Column widths (characters) — 22 columns A–V
COL_WIDTHS = [30, 40, 12, 45, 15, 10, 45, 8, 45, 15, 20, 10, 5, 5, 5, 45, 20, 20, 5, 5, 5, 20]

COL_HEADERS = [
    "Deviations", "Cause", "Drawings /\nReferences",
    "Intermediate Consequence", "Scenario ID", "Cons.\nCat.",
    "Scenario Comments / Final Impacts", "PEC",
    "Mitigation (CME/KME)", "Control\nCategory", "CME Name", "Tags",
    "C", "P", "RL",
    "Recommendation", "CME Name", "Responsibility",
    "C", "P", "RL", "Comments",
]


def generate_excel(report: HAZOPReport) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "HAZOP Report"

    # Set column widths
    for col_idx, width in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    current_row = 1

    # ── Page header row 1 ────────────────────────────────────────────────
    ws.merge_cells(f"A{current_row}:V{current_row}")
    ws[f"A{current_row}"] = (
        "Constitution 2025 Hydrocarbon Processing Systems HAZOP Report    "
        "Document No. CS-OMT-OXY-RM-RPT-00008    Revision No. 1    "
        "Release Date: 12/04/2025"
    )
    _style_header_cell(ws[f"A{current_row}"])
    ws.row_dimensions[current_row].height = 20
    current_row += 1

    # ── Page header row 2 ────────────────────────────────────────────────
    ws.merge_cells(f"A{current_row}:V{current_row}")
    ws[f"A{current_row}"] = (
        f"Systems: {report.system}    "
        f"Nodes: {report.node_name}"
    )
    _style_header_cell(ws[f"A{current_row}"], subheader=True)
    ws.row_dimensions[current_row].height = 20
    current_row += 1

    # ── Column headers ────────────────────────────────────────────────────
    for col_idx, header in enumerate(COL_HEADERS, 1):
        cell = ws.cell(row=current_row, column=col_idx, value=header)
        cell.fill = SUBHEAD_FILL
        cell.font = Font(bold=True, color="FFFFFF", size=9)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER
    ws.row_dimensions[current_row].height = 30
    current_row += 1

    # ── Group deviations by type, ordered ────────────────────────────────
    by_type: dict[str, list[Deviation]] = defaultdict(list)
    for dev in report.deviations:
        by_type[dev.deviation].append(dev)

    deviation_number = 0
    for dev_type in DEVIATION_ORDER:
        devs = by_type.get(dev_type, [])
        if not devs:
            continue
        deviation_number += 1

        dev_block_start = current_row

        # Flatten to (cause_text, deviation_obj) pairs
        all_causes: list[tuple[str, Deviation]] = []
        for dev in devs:
            for cause in (dev.causes or ["No causes identified."]):
                all_causes.append((cause, dev))

        cause_number = 0
        for cause_text, dev in all_causes:
            cause_number += 1
            cause_block_start = current_row
            safeguards = dev.safeguards or []
            risk = dev.risk

            categories = [
                ("PAF",    risk.paf    if risk else None),
                ("PD / LOR", risk.pd_lor if risk else None),
                ("ECR",    risk.ecr    if risk else None),
            ]

            for cat_idx, (cat_label, risk_score) in enumerate(categories):
                cat_block_start = current_row
                rows_in_cat = max(1, len(safeguards))

                for sg_idx in range(rows_in_cat):
                    sg = safeguards[sg_idx] if sg_idx < len(safeguards) else None
                    is_first_sg  = (sg_idx == 0)
                    is_first_cat = (cat_idx == 0)

                    row_data = [
                        # A — filled for first row of first cause; merged across all causes later
                        f"{deviation_number}.  {dev_type}" if (is_first_sg and is_first_cat and cause_number == 1) else "",
                        # B — cause text, first row only; merged across 3 cat blocks later
                        f"{cause_number}.  {cause_text}" if (is_first_sg and is_first_cat) else "",
                        # C — drawing ref
                        (dev.drawing_references[0] if dev.drawing_references else "") if (is_first_sg and is_first_cat) else "",
                        # D — intermediate consequences
                        "\n".join(dev.intermediate_consequences) if (is_first_sg and is_first_cat) else "",
                        # E — scenario ID
                        (dev.deviation_id or "") if (is_first_sg and is_first_cat) else "",
                        # F — cons. cat, first row of each cat block; merged per block later
                        cat_label if is_first_sg else "",
                        # G — scenario comment for this category
                        _scenario_comment_for_cat(dev, cat_label) if is_first_sg else "",
                        # H — PEC
                        (dev.pec or "") if is_first_sg else "",
                        # I — safeguard description
                        sg.description if sg else "",
                        # J — control category
                        sg.pr_classification if sg else "",
                        # K — CME name
                        (sg.cme_name or sg.instrument_tag) if sg else "",
                        # L — tags / PID reference
                        (sg.pid_reference or "") if sg else "",
                        # M — current risk C
                        str(risk_score.consequence) if (risk_score and is_first_sg) else "",
                        # N — current risk P
                        str(risk_score.probability) if (risk_score and is_first_sg) else "",
                        # O — current risk RL
                        str(risk_score.risk_level) if (risk_score and is_first_sg) else "",
                        # P — recommendations
                        "\n".join(dev.recommendations) if (dev.recommendations and is_first_sg and is_first_cat) else "",
                        # Q — planned CME name (empty)
                        "",
                        # R — responsibility
                        (dev.responsibility or "") if (is_first_sg and is_first_cat) else "",
                        # S — planned risk C
                        _planned_c(dev, cat_label) if is_first_sg else "",
                        # T — planned risk P
                        _planned_p(dev, cat_label) if is_first_sg else "",
                        # U — planned risk RL
                        _planned_rl(dev, cat_label) if is_first_sg else "",
                        # V — comments (empty)
                        "",
                    ]

                    for col_idx, value in enumerate(row_data, 1):
                        cell = ws.cell(row=current_row, column=col_idx, value=value)
                        cell.alignment = Alignment(vertical="top", wrap_text=True)
                        cell.border = THIN_BORDER
                        cell.font = Font(size=9)
                    ws.row_dimensions[current_row].height = 40
                    current_row += 1

                # Merge Cons. Cat. (col F) across safeguard rows within this cat block
                if rows_in_cat > 1:
                    ws.merge_cells(f"F{cat_block_start}:F{current_row - 1}")
                    ws[f"F{cat_block_start}"].alignment = Alignment(
                        horizontal="center", vertical="top", wrap_text=True
                    )

            # Merge cause-level columns (B, D, E, G, H, P, R) across all 3 cat blocks
            cause_rows = current_row - cause_block_start
            if cause_rows > 1:
                for col_letter in ("B", "D", "E", "G", "H", "P", "R"):
                    ws.merge_cells(f"{col_letter}{cause_block_start}:{col_letter}{current_row - 1}")
                    ws[f"{col_letter}{cause_block_start}"].alignment = Alignment(
                        vertical="top", wrap_text=True
                    )

        # Merge Deviation (col A) across all cause rows
        dev_rows = current_row - dev_block_start
        if dev_rows > 1:
            ws.merge_cells(f"A{dev_block_start}:A{current_row - 1}")
        cell = ws[f"A{dev_block_start}"]
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.font = Font(bold=True, size=9)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


# ── Helpers ───────────────────────────────────────────────────────────────

def _style_header_cell(cell, subheader: bool = False):
    cell.fill = SUBHEAD_FILL if subheader else HEADER_FILL
    cell.font = Font(bold=True, color="FFFFFF", size=10 if subheader else 11)
    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    cell.border = THIN_BORDER


PAF_KEYWORDS   = {"personnel", "pii", "injury", "fire", "explosion", "vce", "jet", "fatality", "pec"}
PDLOR_KEYWORDS = {"downtime", "production", "repair", "damage", "curtail", "weeks", "months"}
ECR_KEYWORDS   = {"environmental", "spill", "release", "remediation", "ecr", "bbl"}


def _scenario_comment_for_cat(dev: Deviation, cat: str) -> str:
    """
    Assign consequence text to PAF / PD/LOR / ECR based on keywords.
    Falls back to deviation.scenario_comments for PAF.
    """
    buckets: dict[str, list[str]] = {"PAF": [], "PD / LOR": [], "ECR": []}
    for line in dev.consequences:
        lower = line.lower()
        if any(k in lower for k in ECR_KEYWORDS):
            buckets["ECR"].append(line)
        elif any(k in lower for k in PDLOR_KEYWORDS):
            buckets["PD / LOR"].append(line)
        else:
            buckets["PAF"].append(line)

    result = "\n".join(buckets.get(cat, []))
    if not result and cat == "PAF":
        result = dev.scenario_comments or ""
    return result


def _planned_c(dev: Deviation, cat: str) -> str:
    r = dev.planned_residual_risk
    if not r:
        return ""
    score = r.paf if cat == "PAF" else (r.pd_lor if "LOR" in cat else r.ecr)
    return str(score.consequence) if score else ""


def _planned_p(dev: Deviation, cat: str) -> str:
    r = dev.planned_residual_risk
    if not r:
        return ""
    score = r.paf if cat == "PAF" else (r.pd_lor if "LOR" in cat else r.ecr)
    return str(score.probability) if score else ""


def _planned_rl(dev: Deviation, cat: str) -> str:
    r = dev.planned_residual_risk
    if not r:
        return ""
    score = r.paf if cat == "PAF" else (r.pd_lor if "LOR" in cat else r.ecr)
    return str(score.risk_level) if score else ""

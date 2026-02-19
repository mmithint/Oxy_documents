"""
Document Intelligence Service — Knowledge Document Structured Extraction

Extracts text and tables from knowledge documents (PDF, DOCX) in reading order.
Tables are formatted as markdown for better embedding quality.
Text and tables are interleaved by position (Y-axis for PDF, span offset for DOCX)
so that context immediately surrounding a table is preserved in the right order.

Flow:
  Knowledge Doc → Azure Document Intelligence (prebuilt-layout)
    → Extract tables as markdown grids
    → Extract text lines/paragraphs, skipping regions covered by tables
    → Sort text + tables by position
    → Return list of per-page dicts: [{page: int, content: str}, ...]

Returns:
  list[dict] — one entry per page (PDF) or one entry for the whole doc (DOCX)
  Each dict: {"page": int, "content": str}
"""

import asyncio
import io
from collections import defaultdict

import openpyxl
from azure.ai.documentintelligence import DocumentIntelligenceClient
from azure.core.credentials import AzureKeyCredential

from app.core.config import get_settings

settings = get_settings()


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _table_to_markdown(grid: list[list[str]]) -> str:
    """Convert a row/column grid to a GitHub-flavoured markdown table."""
    if not grid:
        return ""

    cols = max(len(row) for row in grid)

    def fmt_row(row: list[str]) -> str:
        padded = list(row) + [""] * (cols - len(row))
        return "| " + " | ".join(str(c).replace("|", "\\|") for c in padded) + " |"

    lines = [fmt_row(grid[0]), "| " + " | ".join(["---"] * cols) + " |"]
    for row in grid[1:]:
        lines.append(fmt_row(row))

    return "\n".join(lines)


def _polygons_overlap(poly_a: list, poly_b: list) -> bool:
    """
    Axis-aligned bounding-box overlap check for two Document Intelligence polygons.
    Polygon format: [x1, y1, x2, y2, x3, y3, x4, y4]
    """
    if not poly_a or len(poly_a) < 8 or not poly_b or len(poly_b) < 8:
        return False

    ax = [poly_a[i] for i in range(0, 8, 2)]
    ay = [poly_a[i] for i in range(1, 8, 2)]
    bx = [poly_b[i] for i in range(0, 8, 2)]
    by = [poly_b[i] for i in range(1, 8, 2)]

    return (
        min(ax) < max(bx) and max(ax) > min(bx)
        and min(ay) < max(by) and max(ay) > min(by)
    )


# ---------------------------------------------------------------------------
# Excel extraction (runs in a thread)
# ---------------------------------------------------------------------------

def _extract_excel_sync(file_bytes: bytes) -> list[dict]:
    """
    Extract structured content from an Excel workbook (.xlsx).

    Each worksheet becomes one page dict:
        {"page": sheet_index, "content": "## Sheet Name\\n\\n[TABLE]...markdown...[/TABLE]"}

    Sheet name is used as a heading so the chunker and LLM know which
    tab the content came from (e.g., "## Risk Matrix", "## Consequence Table").
    Empty rows are skipped. No Azure DI call is made — openpyxl reads the
    native cell values directly, so structure is always perfect.
    """
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    pages: list[dict] = []

    for sheet_idx, sheet_name in enumerate(wb.sheetnames, start=1):
        ws = wb[sheet_name]

        rows: list[list[str]] = []
        for row in ws.iter_rows(values_only=True):
            str_cells = [str(c) if c is not None else "" for c in row]
            # Skip rows that are entirely empty
            if any(c.strip() for c in str_cells):
                rows.append(str_cells)

        if not rows:
            continue

        markdown = _table_to_markdown(rows)
        content = f"## {sheet_name}\n\n[TABLE]\n{markdown}\n[/TABLE]\n"
        pages.append({"page": sheet_idx, "content": content})

    wb.close()
    return pages


# ---------------------------------------------------------------------------
# Synchronous extraction (runs in a thread)
# ---------------------------------------------------------------------------

def _extract_structured_sync(file_bytes: bytes, filename: str) -> list[dict]:
    """
    Blocking extraction — called via asyncio.to_thread to avoid blocking FastAPI.

    Returns:
        list[dict]: [{"page": int, "content": str}, ...]
        For DOCX the list contains a single entry with page=1.
    """
    is_docx = filename.lower().endswith(".docx")
    content_type = (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        if is_docx
        else "application/pdf"
    )

    client = DocumentIntelligenceClient(
        endpoint=settings.AZURE_DOC_INTELLIGENCE_ENDPOINT,
        credential=AzureKeyCredential(settings.AZURE_DOC_INTELLIGENCE_KEY),
    )

    poller = client.begin_analyze_document(
        "prebuilt-layout",
        file_bytes,
        content_type=content_type,
    )
    result = poller.result()

    if not result:
        return []

    # -----------------------------------------------------------------------
    # Step 1 — Extract tables
    # -----------------------------------------------------------------------
    tables_by_page: dict[int, list[dict]] = defaultdict(list)
    table_regions_by_page: dict[int, list] = defaultdict(list)
    table_cell_contents: set[str] = set()

    if hasattr(result, "tables") and result.tables:
        for t_idx, table in enumerate(result.tables):
            rows_count = table.row_count
            cols_count = table.column_count

            # Build grid — handle row/column spans
            grid: list[list[str]] = [[""] * cols_count for _ in range(rows_count)]
            for cell in table.cells:
                r, c = cell.row_index, cell.column_index
                rs = cell.row_span or 1
                cs = cell.column_span or 1
                cell_text = (cell.content or "").strip()
                if cell_text:
                    table_cell_contents.add(cell_text)
                for dr in range(rs):
                    for dc in range(cs):
                        rr, cc = r + dr, c + dc
                        if rr < rows_count and cc < cols_count:
                            grid[rr][cc] = cell_text

            markdown = _table_to_markdown(grid)

            span_offset = (
                table.spans[0].offset
                if hasattr(table, "spans") and table.spans
                else t_idx * 1000
            )

            if table.bounding_regions:
                br = table.bounding_regions[0]
                page_number = br.page_number
                y_pos = br.polygon[1] if br.polygon else 0
                tables_by_page[page_number].append(
                    {"markdown": markdown, "y_position": y_pos, "span_offset": span_offset}
                )
                table_regions_by_page[page_number].append(br)
            else:
                # DOCX tables often lack bounding regions
                tables_by_page[1].append(
                    {"markdown": markdown, "y_position": 0, "span_offset": span_offset}
                )

    # -----------------------------------------------------------------------
    # Step 2a — PDF: extract per-page lines, skip table-covered regions
    # -----------------------------------------------------------------------
    page_docs: list[dict] = []

    if not is_docx:
        if hasattr(result, "pages") and result.pages:
            for page in result.pages:
                pn = page.page_number
                table_regions = table_regions_by_page.get(pn, [])
                all_items: list[dict] = []

                if page.lines:
                    for line in page.lines:
                        line_poly = line.polygon if hasattr(line, "polygon") else None
                        overlaps = any(
                            _polygons_overlap(line_poly, tr.polygon)
                            for tr in table_regions
                            if hasattr(tr, "polygon")
                        )
                        if overlaps:
                            continue
                        y_pos = line_poly[1] if line_poly and len(line_poly) > 1 else 0
                        all_items.append(
                            {"type": "text", "content": line.content, "y_position": y_pos}
                        )

                for t in tables_by_page.get(pn, []):
                    all_items.append(
                        {
                            "type": "table",
                            "content": f"\n[TABLE]\n{t['markdown']}\n[/TABLE]\n",
                            "y_position": t["y_position"],
                        }
                    )

                all_items.sort(key=lambda x: x["y_position"])
                page_content = "\n".join(item["content"] for item in all_items)

                if page_content.strip():
                    page_docs.append({"page": pn, "content": page_content})

    # -----------------------------------------------------------------------
    # Step 2b — DOCX: paragraph-based, ordered by span offset
    # -----------------------------------------------------------------------
    else:
        if hasattr(result, "paragraphs") and result.paragraphs:
            all_table_regions = [
                r for regions in table_regions_by_page.values() for r in regions
            ]
            all_items = []

            for para_idx, para in enumerate(result.paragraphs):
                text = (getattr(para, "content", None) or "").strip()
                if not text:
                    continue
                # Skip text that is actually table cell content
                if text in table_cell_contents:
                    continue

                # Skip paragraphs spatially covered by a table
                if hasattr(para, "bounding_regions") and para.bounding_regions:
                    pr = para.bounding_regions[0]
                    pr_poly = pr.polygon if hasattr(pr, "polygon") else None
                    if any(
                        _polygons_overlap(pr_poly, tr.polygon)
                        for tr in all_table_regions
                        if hasattr(tr, "polygon")
                    ):
                        continue

                # Promote headings with a markdown prefix
                role = (getattr(para, "role", None) or "").lower()
                if role in ("title", "sectionheading", "heading"):
                    text = f"\n## {text}\n"

                span_offset = (
                    para.spans[0].offset
                    if hasattr(para, "spans") and para.spans
                    else para_idx * 100
                )
                all_items.append(
                    {"type": "text", "content": text, "span_offset": span_offset}
                )

            # Add tables in document order
            for t_list in tables_by_page.values():
                for t in t_list:
                    all_items.append(
                        {
                            "type": "table",
                            "content": f"\n[TABLE]\n{t['markdown']}\n[/TABLE]\n",
                            "span_offset": t["span_offset"],
                        }
                    )

            all_items.sort(key=lambda x: x["span_offset"])
            full_content = "\n".join(item["content"] for item in all_items)

            if full_content.strip():
                page_docs.append({"page": 1, "content": full_content})

    return page_docs


# ---------------------------------------------------------------------------
# Public service class
# ---------------------------------------------------------------------------

class KnowledgeDocIntelligenceService:
    """
    Extracts structured text and tables from knowledge documents (PDF, DOCX).

    Returns one dict per page:
        {"page": int, "content": str}

    Tables are formatted as markdown so the embedding model can understand them.
    The Azure SDK call runs in a thread executor to avoid blocking FastAPI's
    event loop.
    """

    async def extract_pages(self, file_content: bytes, filename: str) -> list[dict]:
        """
        Extract structured content from a knowledge document.

        Supported formats:
          - PDF  → Azure Document Intelligence (OCR + layout)
          - DOCX → Azure Document Intelligence (paragraph + table extraction)
          - XLSX → openpyxl (native cell values, no OCR needed)

        Args:
            file_content: Raw bytes of the uploaded file.
            filename:     Original filename — used to detect file type.

        Returns:
            list[dict]: [{"page": int, "content": str}, ...]
        """
        print(f"[Knowledge OCR] Extracting: {filename}")
        try:
            if filename.lower().endswith(".xlsx"):
                # Excel: bypass Azure DI entirely — openpyxl reads native cell values
                pages = await asyncio.to_thread(_extract_excel_sync, file_content)
            else:
                # PDF / DOCX: Azure Document Intelligence
                pages = await asyncio.to_thread(
                    _extract_structured_sync, file_content, filename
                )

            total_chars = sum(len(p["content"]) for p in pages)
            print(f"[Knowledge OCR] {len(pages)} page(s), {total_chars} chars total")
            return pages
        except Exception as e:
            print(f"[Knowledge OCR] Extraction failed: {e}")
            return []


# Module-level singleton
knowledge_doc_intelligence = KnowledgeDocIntelligenceService()
